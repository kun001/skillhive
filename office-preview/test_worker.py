import datetime
import tempfile
import unittest
from pathlib import Path

import openpyxl
from worker import spreadsheet, display


class SpreadsheetTests(unittest.TestCase):
    def test_limits_hidden_sheets_merges_and_saved_values(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "test.xlsx"
            book = openpyxl.Workbook()
            sheet = book.active
            sheet.title = "数据"
            sheet.merge_cells("A1:C1")
            sheet["A1"] = "合并标题"
            sheet["A2"] = datetime.date(2026, 10, 3)
            sheet["B2"] = 0.125
            sheet["B2"].number_format = "0.0%"
            sheet["C2"] = "=1+2"  # no cached value; never evaluated
            sheet.cell(101, 21, "OUTSIDE_PREVIEW")
            hidden = book.create_sheet("隐藏")
            hidden.sheet_state = "hidden"
            for i in range(4):
                book.create_sheet(f"表{i}")["A1"] = f"工作表{i}"
            book.save(source)
            result = spreadsheet(source)
            self.assertEqual(len(result["sheets"]), 3)
            self.assertNotIn("隐藏", [s["name"] for s in result["sheets"]])
            preview = result["sheets"][0]
            self.assertEqual(len(preview["rows"]), 100)
            self.assertEqual(len(preview["rows"][0]), 20)
            self.assertTrue(preview["truncated"])
            self.assertEqual(preview["rows"][1][0]["text"], "2026-10-03")
            self.assertEqual(preview["rows"][1][1]["text"], "12.5%")
            self.assertEqual(preview["rows"][1][2]["text"], "")
            self.assertEqual(preview["merges"], [{"row": 0, "column": 0, "rowSpan": 1, "columnSpan": 3}])

    def test_numbers(self):
        self.assertEqual(display(1234.5, '"￥"#,##0.00'), "￥1,234.50")
        self.assertEqual(display(True), "TRUE")
        self.assertEqual(display("<script>alert(1)</script>"), "<script>alert(1)</script>")


if __name__ == "__main__":
    unittest.main()
